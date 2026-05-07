#!/usr/bin/env python3
"""
CityCred - Cerebro Hibrido Profesional
Backend Flask - Versión para PythonAnywhere
"""

from __future__ import annotations

import json
import os
import re
import unicodedata
from datetime import datetime
from pathlib import Path
from typing import Any

from flask import Flask, jsonify, request, send_from_directory
from flask_cors import CORS

try:
    from openai import OpenAI
except Exception:
    OpenAI = None

try:
    import docx
except Exception:
    docx = None

try:
    import openpyxl
except Exception:
    openpyxl = None


# En PythonAnywhere, el archivo está en /home/santiago123/mysite/
BASE_DIR = Path(__file__).resolve().parent
FRONTEND_DIR = BASE_DIR / "frontend"
DATA_DIR = BASE_DIR / "data"
UPLOADS_DIR = DATA_DIR / "uploads"
SESSION_MEMORY_FILE = DATA_DIR / "session_memory.json"
KNOWLEDGE_FILE = DATA_DIR / "knowledge_memory.json"

DATA_DIR.mkdir(parents=True, exist_ok=True)
UPLOADS_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__, static_folder=str(BASE_DIR), template_folder=str(FRONTEND_DIR))
CORS(app)

OPENAI_RUNTIME_KEY: str | None = None
OPENAI_MODEL = os.getenv("OPENAI_MODEL", "gpt-4o-mini")


def now_iso() -> str:
    return datetime.now().isoformat()


def norm(text: str | None) -> str:
    if not text:
        return ""
    t = text.lower()
    t = unicodedata.normalize("NFD", t)
    t = t.encode("ascii", "ignore").decode("utf-8")
    t = re.sub(r"[^\w\s]", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def write_json(path: Path, data: Any) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def get_openai_key() -> str | None:
    return OPENAI_RUNTIME_KEY or os.getenv("OPENAI_API_KEY")


def detect_force(text: str) -> str:
    t = norm(text)
    if any(x in t for x in ["ejercito", "arma de tierra"]):
        return "Ejército"
    if any(x in t for x in ["armada", "marina"]):
        return "Armada"
    if any(x in t for x in ["gendarmeria", "gna"]):
        return "Gendarmería"
    if any(x in t for x in ["prefectura", "pna"]):
        return "Prefectura"
    if any(x in t for x in ["fuerza aerea", "faa", "aerea"]):
        return "Fuerza Aérea"
    return "General"


def detect_type(text: str) -> str | None:
    t = norm(text)
    if any(x in t for x in ["voluntario", "voluntaria", "tropa"]):
        return "voluntario"
    if any(x in t for x in ["carrera", "cuadro permanente"]):
        return "carrera"
    return None


def detect_cupo(text: str) -> str | None:
    t = norm(text)
    if not any(x in t for x in ["cupo", "afectacion", "disponible"]):
        return None
    m = re.search(r"(\d{4,8})", t)
    return m.group(1) if m else "mencionado"


def detect_intent(message: str, history: str) -> tuple[str, float]:
    t = norm(message)
    h = norm(history)
    scores = {
        "inicio": 0,
        "desconfianza": 0,
        "requisitos": 0,
        "explicar_cupo": 0,
        "guia_cupo": 0,
        "tipo": 0,
        "cupo": 0,
        "avanzar": 0,
        "duda": 0,
        "cierre": 0,
        "general": 1,
    }
    kws = {
        "inicio": ["hola", "buenas", "quiero info", "info"],
        "desconfianza": ["estafa", "no confio", "es seguro", "son reales", "miedo"],
        "requisitos": ["requisito", "dni", "recibo", "cbu", "documento"],
        "explicar_cupo": ["que es el cupo", "que es cupo", "explica cupo"],
        "guia_cupo": ["como saco", "donde saco", "guiame", "como pido"],
        "tipo": ["voluntario", "voluntaria", "carrera", "tropa"],
        "cupo": ["mi cupo", "afectacion", "disponible", "cupo"],
        "avanzar": ["avanzar", "me interesa", "dale", "quiero sacar", "quiero seguir"],
        "duda": ["no se", "nose", "duda", "mmm"],
        "cierre": ["listo", "ok", "mandame", "cerrar", "hagamoslo"],
    }
    for intent, words in kws.items():
        for w in words:
            if w in t:
                scores[intent] += 3
            if w in h:
                scores[intent] += 1
    best = max(scores.items(), key=lambda x: x[1])
    conf = min(0.99, max(0.35, best[1] / 12.0))
    return best[0], round(conf, 2)


def stage(force: str, person_type: str | None, cupo: str | None, intent: str) -> str:
    if force == "General":
        return "paso_2_fuerza"
    if not person_type:
        return "paso_3_tipo"
    if not cupo:
        return "paso_4_5_cupo"
    if intent in ["requisitos", "avanzar"]:
        return "paso_6_7_datos"
    if intent == "cierre":
        return "paso_8_cierre"
    return "paso_6_montos"


def cupo_guide(force: str) -> str:
    if force == "Ejército":
        return "Haberes 2.0 -> Prestamos -> Nueva solicitud."
    if force == "Armada":
        return "Portal de haberes Armada -> Monto de afectacion."
    if force == "Gendarmería":
        return "UTAC o Administracion indicando AMFAYS codigo 400571."
    if force == "Prefectura":
        return "En recibo: MONTO DISPONIBLE DE DEDUCCION."
    if force == "Fuerza Aérea":
        return "Contaduria/Administracion indicando AMFAYS codigo 400571."
    return "Decime tu fuerza y te guio exacto."


def missing(
    force: str, person_type: str | None, cupo: str | None, full_text: str
) -> list[str]:
    t = norm(full_text)
    out: list[str] = []
    if force == "General":
        out.append("fuerza")
    if force != "General" and not person_type:
        out.append("tipo (voluntario/carrera)")
    if not cupo:
        out.append("cupo")
    if "dni" not in t:
        out.append("dni")
    if "recibo" not in t:
        out.append("recibo")
    if "mail" not in t:
        out.append("mail")
    if "cbu" not in t:
        out.append("cbu")
    return out


def score_interest(intent: str, full_text: str) -> int:
    t = norm(full_text)
    s = 40
    if intent in ["avanzar", "cierre"]:
        s += 35
    if intent == "cupo":
        s += 18
    if any(x in t for x in ["quiero", "me interesa", "dale"]):
        s += 12
    return max(0, min(100, s))


def score_distrust(intent: str, full_text: str) -> int:
    t = norm(full_text)
    s = 10
    if intent == "desconfianza":
        s += 60
    for kw in ["estafa", "miedo", "no confio", "duda"]:
        if kw in t:
            s += 10
    return max(0, min(100, s))


def score_close(stage_name: str, interest: int, distrust: int, miss: list[str]) -> int:
    p = interest - int(distrust * 0.6)
    if stage_name == "paso_6_7_datos":
        p += 20
    if stage_name == "paso_8_cierre":
        p += 30
    p -= min(30, len(miss) * 6)
    return max(0, min(100, p))


def next_action(stage_name: str, miss: list[str], intent: str) -> str:
    if "fuerza" in miss:
        return "preguntar_fuerza"
    if any("tipo" in m for m in miss):
        return "preguntar_tipo"
    if "cupo" in miss:
        return "pedir_cupo"
    if intent in ["requisitos", "avanzar"]:
        return "pedir_documentacion"
    if stage_name in ["paso_6_montos", "paso_6_7_datos"]:
        return "cerrar_operacion"
    return "seguir_conversacion"


def vendedor_guide() -> str:
    return (
        "Flujo obligatorio:\n"
        "1) Saludo\n2) Fuerza\n3) Voluntario/carrera\n4) Pedir cupo\n"
        "5) Explicar como sacarlo\n6) Armar montos\n7) Pedir datos\n8) Cerrar"
    )


def local_answer(
    force: str, person_type: str | None, intent: str, miss: list[str]
) -> str:
    if intent == "inicio" and force == "General":
        return "Hola 👋 trabajamos con créditos por Decreto 14/12 con descuento por haberes. ¿De qué fuerza sos?"
    if intent == "inicio" and force != "General" and not person_type:
        return (
            "Perfecto. Para seguir, confirmame: ¿sos personal voluntario o de carrera?"
        )
    if intent == "inicio" and force != "General" and person_type and "cupo" in miss:
        return f"Perfecto. Ya tengo fuerza y tipo. Ahora necesito tu cupo. {cupo_guide(force)}"
    if intent == "desconfianza":
        return "Te entiendo 👍 Somos Citycred, trabajamos con AMFAYS y tenemos oficinas en Río Negro y Bahía Blanca."
    if force != "General" and not person_type:
        return "Perfecto. ¿sos personal voluntario o de carrera?"
    if "cupo" in miss:
        return f"Perfecto. Para avanzar necesito tu cupo. {cupo_guide(force)}"
    if intent in ["requisitos", "avanzar"]:
        return "Para avanzar pasame: DNI, último recibo, cupo, mail y CBU Banco Nación."
    return "Perfecto, vamos bien. Confirmame tus datos y te armo la mejor opción para cerrar hoy."


def load_sessions() -> dict[str, Any]:
    return read_json(SESSION_MEMORY_FILE, {"sessions": {}})


def save_sessions(data: dict[str, Any]) -> None:
    write_json(SESSION_MEMORY_FILE, data)


def update_session(
    session_id: str, payload: dict[str, Any], analysis: dict[str, Any]
) -> dict[str, Any]:
    db = load_sessions()
    sessions = db.setdefault("sessions", {})
    s = sessions.setdefault(
        session_id,
        {"created_at": now_iso(), "messages": [], "facts": {}, "last": {}},
    )
    s["messages"].append(
        {
            "ts": now_iso(),
            "mensaje": payload.get("mensaje") or payload.get("ultimo_mensaje") or "",
            "historial_len": len(payload.get("historial", "") or ""),
        }
    )
    s["facts"].update(
        {
            "fuerza": analysis.get("fuerza_detectada"),
            "tipo": analysis.get("tipo_detectado"),
            "cupo": analysis.get("cupo_detectado"),
            "stage": analysis.get("etapa_comercial"),
        }
    )
    s["last"] = analysis
    save_sessions(db)
    return s


def load_knowledge() -> dict[str, Any]:
    return read_json(KNOWLEDGE_FILE, {"items": [], "updated_at": now_iso()})


def save_knowledge(k: dict[str, Any]) -> None:
    k["updated_at"] = now_iso()
    write_json(KNOWLEDGE_FILE, k)


def extract_file_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".txt":
        return path.read_text(encoding="utf-8", errors="ignore")
    if ext == ".docx" and docx is not None:
        d = docx.Document(str(path))
        return "\n".join(p.text for p in d.paragraphs if p.text.strip())
    if ext == ".xlsx" and openpyxl is not None:
        wb = openpyxl.load_workbook(path, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets:
            lines.append(f"[Hoja {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if vals:
                    lines.append(" | ".join(vals))
        return "\n".join(lines)
    return ""


def knowledge_lookup(message: str) -> str | None:
    items = load_knowledge().get("items", [])
    if not items:
        return None
    m = norm(message)
    for item in reversed(items):
        txt = norm(item.get("content", ""))
        if not txt:
            continue
        for token in m.split()[:5]:
            if token and token in txt:
                raw = item.get("content", "")
                if len(raw.strip()) > 50:
                    return raw[:900]
    return None


def openai_answer(
    payload: dict[str, Any], analysis: dict[str, Any], base_text: str
) -> str | None:
    key = get_openai_key()
    if not key or OpenAI is None:
        return None

    client = OpenAI(api_key=key)
    system_prompt = (
        "Sos vendedor senior de CityCred. Responde natural, breve, comercial y humano. "
        "Respeta SIEMPRE este flujo: saludo, fuerza, tipo, cupo, montos, datos, cierre. "
        "Regla inviolable: si hay fuerza detectada y falta tipo, preguntar textual: "
        "'¿sos personal voluntario o de carrera?'. No repetir datos ya confirmados."
    )
    user_payload = {
        "mensaje": payload.get("mensaje") or payload.get("ultimo_mensaje") or "",
        "historial": payload.get("historial", ""),
        "datos_conocidos": payload.get("datos_conocidos", {}),
        "analisis": analysis,
        "respuesta_base": base_text,
    }

    try:
        resp = client.chat.completions.create(
            model=OPENAI_MODEL,
            temperature=0.3,
            messages=[
                {"role": "system", "content": system_prompt},
                {
                    "role": "user",
                    "content": json.dumps(user_payload, ensure_ascii=False),
                },
            ],
        )
        text = (resp.choices[0].message.content or "").strip()
        return text if text else None
    except Exception:
        return None


def enforce_rules(answer: str, analysis: dict[str, Any]) -> str:
    out = answer.strip()
    force = analysis.get("fuerza_detectada", "General")
    person_type = analysis.get("tipo_detectado")
    if (
        force in ["Ejército", "Armada", "Gendarmería", "Prefectura", "Fuerza Aérea"]
        and not person_type
        and "voluntario" not in norm(out)
        and "carrera" not in norm(out)
    ):
        out += "\n\n¿sos personal voluntario o de carrera?"
    return out


def analyze(payload: dict[str, Any]) -> dict[str, Any]:
    msg = payload.get("mensaje") or payload.get("ultimo_mensaje") or ""
    hist = payload.get("historial", "") or ""
    known = payload.get("datos_conocidos", {}) or {}
    full = f"{hist}\n{msg}"

    force = detect_force(full)
    if force == "General":
        force = payload.get("fuerza", "General")
    if force == "General" and isinstance(known, dict):
        force = known.get("fuerza", "General")

    person_type = detect_type(full) or payload.get("tipo")
    if not person_type and isinstance(known, dict):
        person_type = known.get("tipo")

    cupo = detect_cupo(full)
    if not cupo and isinstance(known, dict):
        cupo = known.get("cupo")

    intent, confidence = detect_intent(msg, hist)
    stage_name = stage(force, person_type, cupo, intent)
    miss = missing(force, person_type, cupo, full)
    interest = score_interest(intent, full)
    distrust = score_distrust(intent, full)
    close_prob = score_close(stage_name, interest, distrust, miss)
    client_state = "caliente" if (interest >= 70 and distrust < 45) else "frio"
    action = next_action(stage_name, miss, intent)
    objeciones: list[str] = []
    txt = norm(full)
    if any(
        x in txt for x in ["estafa", "no confio", "es seguro", "son reales", "miedo"]
    ):
        objeciones.append("desconfianza")
    if "no se" in txt or "duda" in txt:
        objeciones.append("duda")

    return {
        "intencion_detectada": intent,
        "nivel_confianza": confidence,
        "fuerza_detectada": force,
        "tipo_detectado": person_type,
        "cupo_detectado": cupo,
        "etapa_comercial": stage_name,
        "datos_faltantes": miss,
        "nivel_interes": interest,
        "nivel_desconfianza": distrust,
        "probabilidad_cierre": close_prob,
        "cliente_estado": client_state,
        "proxima_accion_recomendada": action,
        "objeciones": objeciones,
    }


def build_response(payload: dict[str, Any], prefer_openai: bool) -> dict[str, Any]:
    analysis = analyze(payload)

    base = local_answer(
        analysis["fuerza_detectada"],
        analysis["tipo_detectado"],
        analysis["intencion_detectada"],
        analysis["datos_faltantes"],
    )

    hit = knowledge_lookup(
        payload.get("mensaje") or payload.get("ultimo_mensaje") or ""
    )
    if hit and analysis["intencion_detectada"] in ["general", "duda"]:
        base = f"{base}\n\nReferencia interna:\n{hit[:500]}"

    final = base
    if prefer_openai:
        session_id = payload.get("session_id", "default")
        sessions = load_sessions().get("sessions", {})
        session_memory = sessions.get(session_id, {})
        openai_payload = dict(payload)
        openai_payload["memoria_sesion"] = session_memory
        openai_payload["analisis_comercial"] = {
            "fuerza": analysis["fuerza_detectada"],
            "tipo": analysis["tipo_detectado"],
            "cupo": analysis["cupo_detectado"],
            "etapa_comercial": analysis["etapa_comercial"],
            "intencion": analysis["intencion_detectada"],
            "objeciones": analysis.get("objeciones", []),
            "datos_faltantes": analysis["datos_faltantes"],
        }
        improved = openai_answer(openai_payload, analysis, base)
        if improved:
            final = improved

    final = enforce_rules(final, analysis)

    session_id = payload.get("session_id", "default")
    session = update_session(session_id, payload, analysis)

    return {
        "respuesta_sugerida": final,
        "intencion_detectada": analysis["intencion_detectada"],
        "proxima_accion_recomendada": analysis["proxima_accion_recomendada"],
        "datos_faltantes": analysis["datos_faltantes"],
        "guia_interna": vendedor_guide(),
        "fuerza_detectada": analysis["fuerza_detectada"],
        "tipo_detectado": analysis["tipo_detectado"],
        "nivel_confianza": analysis["nivel_confianza"],
        "etapa_comercial": analysis["etapa_comercial"],
        "nivel_interes": analysis["nivel_interes"],
        "nivel_desconfianza": analysis["nivel_desconfianza"],
        "probabilidad_cierre": analysis["probabilidad_cierre"],
        "cliente_estado": analysis["cliente_estado"],
        "objeciones": analysis.get("objeciones", []),
        "memoria_sesion": {
            "session_id": session_id,
            "mensajes": len(session.get("messages", [])),
            "facts": session.get("facts", {}),
        },
        "openai_activo": bool(get_openai_key()),
        "modelo_openai": OPENAI_MODEL,
        "timestamp": now_iso(),
    }


@app.get("/")
def index() -> Any:
    return send_from_directory(FRONTEND_DIR, "simulador.html")


@app.post("/api/response")
def api_response() -> Any:
    payload = request.get_json(force=True) or {}
    data = build_response(
        payload, prefer_openai=bool(payload.get("usar_openai", False))
    )
    return jsonify(data)


@app.post("/api/chat")
def api_chat() -> Any:
    payload = request.get_json(force=True) or {}
    data = build_response(payload, prefer_openai=True)
    data["endpoint"] = "/api/chat"
    return jsonify(data)


@app.post("/api/upload-info")
def api_upload_info() -> Any:
    if "file" not in request.files:
        return jsonify({"ok": False, "error": "Falta archivo"}), 400
    f = request.files["file"]
    if not f.filename:
        return jsonify({"ok": False, "error": "Archivo invalido"}), 400

    ext = Path(f.filename).suffix.lower()
    if ext not in [".txt", ".docx", ".xlsx"]:
        return jsonify({"ok": False, "error": "Solo .txt .docx .xlsx"}), 400

    filename = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{Path(f.filename).name}"
    out = UPLOADS_DIR / filename
    f.save(out)

    content = extract_file_text(out)
    if not content.strip():
        return jsonify({"ok": False, "error": "No se pudo leer contenido"}), 400

    k = load_knowledge()
    items = k.setdefault("items", [])
    items.append(
        {
            "filename": filename,
            "ext": ext,
            "content": content[:200000],
            "created_at": now_iso(),
        }
    )
    save_knowledge(k)
    return jsonify(
        {
            "ok": True,
            "archivo": filename,
            "items_knowledge": len(items),
            "chars": len(content),
        }
    )


@app.post("/api/openai-status")
def api_openai_status() -> Any:
    return jsonify(
        {"openai_configurado": bool(get_openai_key()), "modelo": OPENAI_MODEL}
    )


@app.post("/api/set-openai-key")
def api_set_openai_key() -> Any:
    global OPENAI_RUNTIME_KEY
    payload = request.get_json(force=True) or {}
    key = (payload.get("api_key") or "").strip()
    temporal = bool(payload.get("temporal", True))
    if not key:
        return jsonify({"ok": False, "error": "api_key vacia"}), 400
    OPENAI_RUNTIME_KEY = key
    if not temporal:
        os.environ["OPENAI_API_KEY"] = key
    return jsonify({"ok": True, "openai_configurado": True, "temporal": temporal})


@app.get("/api/health")
def api_health() -> Any:
    return jsonify(
        {
            "status": "ok",
            "service": "CityCred Cerebro Hibrido",
            "openai_configurado": bool(get_openai_key()),
            "modelo_openai": OPENAI_MODEL,
            "knowledge_items": len(load_knowledge().get("items", [])),
            "sessions": len(load_sessions().get("sessions", {})),
            "timestamp": now_iso(),
        }
    )


@app.route("/<path:filename>")
def serve_static(filename: str) -> Any:
    return send_from_directory(BASE_DIR, filename)
